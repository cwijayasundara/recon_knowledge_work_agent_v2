from __future__ import annotations

import io
import json
from pathlib import Path

import jsonschema
from onboarding_sdk import manifest, review
from onboarding_sdk.entities.affiliate import AffiliateOptions, process
from onboarding_sdk.render import intacct_csv
from openpyxl import load_workbook

from tests.conftest import FIXTURE_DIR
from tests.support.pipeline import canonical_for, expected

BINDINGS = {"affiliate_id": "Affiliate ID", "affiliate_name": "Affiliate Name"}


def _review(name: str, options: AffiliateOptions | None = None, bindings: dict | None = None):  # type: ignore[no-untyped-def]
    spec = expected(name)
    canon = canonical_for(name)
    result = process(canon, options)
    data = review.workbook(
        FIXTURE_DIR / name,
        canon,
        result,
        decisions=[
            {
                "seq": 1,
                "kind": "approve",
                "gate": "brief",
                "actor": "analyst@sponsor-a",
                "at": "2025-01-01T00:00:00Z",
                "payload": {},
            }
        ],
        brief={"source": {"file": name, "sheet": spec["sheet"]}, "questions": []},
        bindings=bindings or spec["bindings"],
        header_row=spec["header_row"],
    )
    return load_workbook(io.BytesIO(data)), data


def test_review_sheets_and_formulas() -> None:
    wb, _ = _review("titled.xlsx")
    assert wb.sheetnames == ["Source", "Upload preview", "Findings", "Decisions", "Brief"]
    source = wb["Source"]
    assert source["A4"].value == "Affiliate ID"
    assert source["B5"].value == "Meridian Capital GP IV, LLC"
    preview = wb["Upload preview"]
    header = [c.value for c in preview[1]]
    assert header[:5] == ["ITEM_ID", "NAME", "ITEM_TYPE", "DESCRIPTION", "DONOTIMPORT"]
    assert preview["B2"].value == "=LEFT(Source!B5,100)"
    assert preview["A2"].value == "=Source!A5"


def test_derived_id_is_value_with_rule_note() -> None:
    wb, _ = _review("ids_missing.csv")
    cell = wb["Upload preview"]["A2"]
    assert cell.value == "MERIDIAN_CAPITAL_GP_IV_LLC"
    assert cell.comment is not None and "affiliate.id.derive_from_name" in cell.comment.text


def test_findings_and_decisions_sheets() -> None:
    wb, _ = _review("edge.csv")
    rows = list(wb["Findings"].iter_rows(values_only=True))
    assert rows[0][:4] == ("code", "severity", "row", "source_row")
    assert rows[1][:4] == ("AFF_ERR_ITEM_ID_BLANK", "error", 1, 2)
    decisions = list(wb["Decisions"].iter_rows(values_only=True))
    assert decisions[1][1] == "approve"


def test_review_bytes_are_deterministic() -> None:
    _, first = _review("clean.csv")
    _, second = _review("clean.csv")
    assert first == second


def test_manifest_validates_against_schema() -> None:
    canon = canonical_for("clean.csv")
    result = process(canon)
    csv_bytes = intacct_csv(result)
    upload = FIXTURE_DIR / "clean.csv"
    doc = manifest.build(
        run_id="run-1",
        sponsor_id="sponsor-a",
        entity="affiliate",
        upload_name=upload.name,
        upload_bytes=upload.read_bytes(),
        recipe={"id": "rcp-1", "sha256": "0" * 64, "origin": "standard", "version": 1},
        options=AffiliateOptions(acknowledged=frozenset({("AFF_WARN_ITEM_ID_DERIVED", 3)})),
        bindings=[
            {"field": "affiliate_id", "column": "Affiliate ID", "route": "ontology_exact"},
            {"field": "affiliate_name", "column": "Affiliate Name", "route": "history"},
        ],
        decisions=[
            {
                "seq": 1,
                "kind": "approve",
                "gate": "brief",
                "actor": "a",
                "at": "2025-01-01T00:00:00Z",
                "payload": {},
            }
        ],
        outputs={"Affiliates.csv": csv_bytes, "review.xlsx": b"xlsx"},
        approvers=[{"gate": "signoff", "actor": "a", "at": "2025-01-01T00:00:00Z"}],
        created_at="2025-01-01T00:00:00Z",
        finalized_at="2025-01-01T00:05:00Z",
        findings=result.findings,
    )
    schema = json.loads(Path(manifest.SCHEMA_PATH).read_text())
    jsonschema.validate(doc, schema)
    assert doc["options"]["acknowledged"] == [["AFF_WARN_ITEM_ID_DERIVED", 3]]
    assert doc["outputs"][0]["name"] == "Affiliates.csv"
    assert len(doc["outputs"][0]["sha256"]) == 64
    json.dumps(doc)


def test_options_round_trip() -> None:
    options = AffiliateOptions(
        item_type="Non-Inventory",
        row_item_types={2: "Inventory"},
        id_overrides={4: "X"},
        excluded_rows={1: "blank"},
        acknowledged=frozenset({("A", None), ("B", 2)}),
    )
    assert AffiliateOptions.from_dict(json.loads(json.dumps(options.to_dict()))) == options
