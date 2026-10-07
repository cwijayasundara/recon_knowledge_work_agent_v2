"""Regression tests for the final-review findings (C1, I1-I8, and the re-graded sponsor/formula issues)."""

from __future__ import annotations

import io
import json
import shlex
from pathlib import Path

import pytest
from onboarding_sdk import read, recipes
from openpyxl import Workbook as XlsxWorkbook
from openpyxl import load_workbook

from onboarding_agent.graph.build import UploadRejected, Workbench
from tests.conftest import FIXTURE_DIR
from tests.support.scripted_model import call, say, tools
from tests.support.scripts import brief_for, report_simple, scope_standard
from tests.support.services import Models, offline_services

A = "analyst@sponsor-a"


def _start(bench: Workbench, name: str, *, file_name: str | None = None, sponsor: str = "sponsor-a") -> str:
    return bench.start(
        sponsor_id=sponsor,
        entity="affiliate",
        file_name=file_name or name,
        data=(FIXTURE_DIR / name).read_bytes(),
        actor=A,
    )


def _finish(bench: Workbench, run_id: str) -> dict:  # type: ignore[type-arg]
    snap = bench.snapshot(run_id)
    for _ in range(8):
        if snap["pending"] is None:
            return snap
        if snap["pending"]["gate"] == "findings":
            acks = [
                {"kind": "acknowledge_finding", "code": f["code"], "row": f["row"]}
                for f in snap["result"]["findings"]
                if f["requires_ack"] and not f["acknowledged"]
            ]
            if acks:
                snap = bench.respond(run_id, {"action": "change", "actor": A, "changes": acks})
                continue
        snap = bench.respond(run_id, {"action": "approve", "actor": A})
    return snap


def test_c1_report_mode_cannot_rebind_or_swap_recipe(tmp_path: Path) -> None:
    # fastpath off: the report-mode guard needs a run where the supervisor scopes and reports.
    models = Models()
    models.supervisor.script = [
        *scope_standard("extra_columns.csv"),
        tools(
            call(
                "write_standard_recipe",
                sheet="extra_columns",
                header_row=1,
                affiliate_id="Affiliate Id",
                affiliate_name="Parent Entity",
            )
        ),
        tools(call("run_pipeline")),
        tools(call("submit_report", report={"summary": "ok"})),
        say("done"),
    ]
    bench = Workbench(offline_services(tmp_path, models, fastpath=False))
    run_id = _start(bench, "extra_columns.csv")
    bench.respond(run_id, {"action": "approve", "actor": A})
    offered = models.supervisor.offered[-1]
    assert "write_standard_recipe" not in offered and "run_pipeline" not in offered
    snap = _finish(bench, run_id)
    assert snap["status"] == "locked"
    csv = bench.services.stores.objects.get(f"runs/{run_id}/outputs/Affiliates.csv").decode()
    assert "Meridian Capital GP IV, LLC" in csv
    assert "Meridian Capital Partners IV" not in csv


def test_c1_spine_uses_the_approved_snapshot_even_if_context_changes(tmp_path: Path) -> None:
    bench = Workbench(offline_services(tmp_path, Models()))
    run_id = _start(bench, "extra_columns.csv")
    bench.respond(run_id, {"action": "approve", "actor": A})
    # Anything agent-writable in the shared context is re-pinned from state.
    ctx = bench.spine._contexts[run_id]
    ctx.bindings = {"affiliate_id": "Affiliate Id", "affiliate_name": "Parent Entity"}
    rogue = tmp_path / "rogue.py"
    rogue.write_text(recipes.standard(ctx.bindings, "extra_columns", 1))
    ctx.candidate_recipe = {"origin": "standard", "path": str(rogue), "sha256": "0" * 64}
    ctx.result = None
    _finish(bench, run_id)
    csv = bench.services.stores.objects.get(f"runs/{run_id}/outputs/Affiliates.csv").decode()
    assert "Meridian Capital GP IV, LLC" in csv and "Partners IV" not in csv
    doc = json.loads(bench.services.stores.objects.get(f"runs/{run_id}/outputs/manifest.json"))
    assert {b["field"]: b["column"] for b in doc["bindings"]}["affiliate_name"] == "Affiliate Name"


def test_c1_recipe_file_changed_after_approval_is_refused(tmp_path: Path) -> None:
    bench = Workbench(offline_services(tmp_path, Models()))
    run_id = _start(bench, "clean.csv")
    snap = bench.respond(run_id, {"action": "approve", "actor": A})
    path = Path(snap["approved"]["recipe"]["path"])
    path.write_text(path.read_text().replace("Affiliate Name", "Fund Complex"))
    bench.spine.forget(run_id)
    snap = bench.respond(
        run_id, {"action": "change", "actor": A, "changes": [{"kind": "set_item_type", "value": "Non-Inventory"}]}
    )
    assert "hash" in snap["gate_message"]


def test_i1_brief_must_match_recipe_at_approval(tmp_path: Path) -> None:
    # fastpath off: the mismatch blocker needs a supervisor-submitted brief.
    models = Models()
    models.supervisor.script = [
        *scope_standard("extra_columns.csv")[:-1],
        # After submitting, the agent re-binds; the brief no longer describes the recipe.
        tools(
            call(
                "propose_bindings",
                sheet="extra_columns",
                header_row=1,
                affiliate_id="Affiliate Id",
                affiliate_name="Parent Entity",
            )
        ),
        say("done"),
    ]
    bench = Workbench(offline_services(tmp_path, models, fastpath=False))
    run_id = _start(bench, "extra_columns.csv")
    snap = bench.snapshot(run_id)
    assert snap["pending"]["blocked_reasons"], snap["pending"]
    snap = bench.respond(run_id, {"action": "approve", "actor": A})
    assert snap["pending"]["gate"] == "brief"
    assert "Cannot approve" in snap["gate_message"]


def test_i1_brief_item_type_is_applied(tmp_path: Path) -> None:
    # fastpath off: the brief with the non-default item type is submitted by the supervisor.
    brief = brief_for("clean.csv") | {"item_type": "Non-Inventory"}
    models = Models()
    models.supervisor.script = [
        *scope_standard("clean.csv")[:3],
        tools(call("submit_brief", brief=brief)),
        say("ok"),
        *report_simple(),
    ]
    bench = Workbench(offline_services(tmp_path, models, fastpath=False))
    run_id = _start(bench, "clean.csv")
    snap = bench.respond(run_id, {"action": "approve", "actor": A})
    codes = {f["code"] for f in snap["result"]["findings"]}
    assert "AFF_WARN_ITEM_TYPE_OVERRIDE" in codes
    assert snap["options"]["item_type"] == "Non-Inventory"


def test_i2_options_reset_when_the_brief_is_approved_again(tmp_path: Path) -> None:
    models = Models()
    # The first brief is the spine's; a layout change re-scopes, and the report follows the agent.
    models.supervisor.script = [
        *scope_standard("ids_missing.csv"),
        *report_simple(),
    ]
    bench = Workbench(offline_services(tmp_path, models))
    run_id = _start(bench, "ids_missing.csv")
    bench.respond(run_id, {"action": "approve", "actor": A})
    bench.respond(
        run_id,
        {
            "action": "change",
            "actor": A,
            "changes": [
                {"kind": "acknowledge_finding", "code": "AFF_WARN_ITEM_ID_DERIVED", "row": 2},
                {"kind": "override_item_id", "row": 4, "value": "X_4"},
            ],
        },
    )
    snap = bench.respond(
        run_id, {"action": "change", "actor": A, "changes": [{"kind": "set_header_row", "header_row": 1}]}
    )
    assert snap["pending"]["gate"] == "brief"
    snap = bench.respond(run_id, {"action": "approve", "actor": A})
    assert snap["options"]["acknowledged"] == [] and snap["options"]["id_overrides"] == {}


def test_i4_upload_names_with_spaces_work_in_the_sandbox(tmp_path: Path) -> None:
    from tests.support.services import context_for

    services = offline_services(tmp_path, Models())
    ctx = context_for(services, "clean.csv")
    target = ctx.upload_path.with_name("Affiliates (1) final.csv")
    ctx.upload_path.rename(target)
    ctx.upload_path = target
    source = recipes.standard({"affiliate_id": "Affiliate ID", "affiliate_name": "Affiliate Name"}, "clean", 1)
    path = tmp_path / "authored.py"
    path.write_text(source)
    from onboarding_sdk.recipes import hashlib as _  # noqa: F401 - module check only

    from onboarding_agent.tools.pipeline import execute_recipe

    canon = execute_recipe(
        ctx, {"origin": "authored", "path": str(path), "sha256": recipes.hashlib.sha256(source.encode()).hexdigest()}
    )
    assert len(canon.rows) == 8
    assert shlex.quote("/in/Affiliates (1) final.csv") in ctx._sandbox.commands[-1]  # type: ignore[union-attr]
    ctx.close()


def test_i5_supervisor_sees_only_its_own_sponsor_notes(tmp_path: Path) -> None:
    models = Models()
    services = offline_services(tmp_path, models)
    other = services.stores.objects.local_path("sponsors/sponsor-b/AGENTS.md")
    other.parent.mkdir(parents=True, exist_ok=True)
    other.write_text("sponsor-b secret layout")
    models.supervisor.script = [
        tools(call("ls", path="/")),
        tools(call("read_file", file_path="/sponsors/sponsor-b/AGENTS.md")),
        say("done"),
    ]
    from onboarding_agent.assembly import invoke_supervisor
    from tests.support.services import context_for

    ctx = context_for(services, "clean.csv")
    result = invoke_supervisor(services, ctx, "scope")
    text = " ".join(str(m.content) for m in result["messages"])
    assert "sponsor-b secret layout" not in text
    system = " ".join(str(m.content) for m in models.supervisor.seen[0])
    assert "sponsor-b" not in system


def test_i6_excel_row_numbers_keep_leading_blank_rows(tmp_path: Path) -> None:
    wb = XlsxWorkbook()
    ws = wb.active
    assert ws is not None
    ws["B3"], ws["C3"] = "Affiliate ID", "Affiliate Name"
    ws["B4"], ws["C4"] = "AFF_1", "Alpha LLC"
    path = tmp_path / "offset.xlsx"
    wb.save(path)
    sheet = read.open(path).sheets[0]
    table = sheet.table(3)
    row = next(table.rows())
    assert row.row_number == 4
    assert row.get("Affiliate Name") == "Alpha LLC"


def test_i7_existing_container_is_replaced_on_start(tmp_path: Path) -> None:
    from onboarding_agent.sandbox.base import SandboxMounts
    from onboarding_agent.sandbox.docker_backend import DockerSandbox

    removed: list[str] = []

    class Old:
        def remove(self, force: bool = False) -> None:
            removed.append("old")

    class Containers:
        def list(self, all: bool, filters: dict) -> list:  # type: ignore[type-arg]
            assert filters == {"label": "onb.run_id=run-1"}
            return [Old()]

        def run(self, **kwargs):  # type: ignore[no-untyped-def]
            return type("C", (), {"id": "new", "remove": lambda self, force=False: None})()

    client = type("Client", (), {"containers": Containers()})()
    for name in ("in", "ref", "skills"):
        (tmp_path / name).mkdir()
    box = DockerSandbox("run-1", SandboxMounts(tmp_path / "in", tmp_path / "ref", tmp_path / "skills"), client=client)
    box.start()
    assert removed == ["old"] and box.id == "new"


def test_i7_rejected_run_releases_its_context(tmp_path: Path) -> None:
    models = Models()
    models.supervisor.script = scope_standard("clean.csv")
    bench = Workbench(offline_services(tmp_path, models))
    run_id = _start(bench, "clean.csv")
    bench.respond(run_id, {"action": "reject", "actor": A, "reason": "wrong file"})
    assert run_id not in bench.spine._contexts


def test_i8_sandbox_commands_ignore_modules_in_work(tmp_path: Path) -> None:
    from onboarding_agent.tools.pipeline import execute_recipe
    from tests.support.services import context_for

    services = offline_services(tmp_path, Models())
    ctx = context_for(services, "clean.csv")
    box = ctx.sandbox()
    fake = b'import json\nprint(json.dumps({"ok": True, "canonical": {"rows": [], "dropped": []}}))\n'
    box.upload_files([("/work/onboarding_sdk/__init__.py", b""), ("/work/onboarding_sdk/recipes.py", fake)])
    source = recipes.standard({"affiliate_id": "Affiliate ID", "affiliate_name": "Affiliate Name"}, "clean", 1)
    path = tmp_path / "authored.py"
    path.write_text(source)
    canon = execute_recipe(
        ctx, {"origin": "authored", "path": str(path), "sha256": recipes.hashlib.sha256(source.encode()).hexdigest()}
    )
    assert len(canon.rows) == 8
    ctx.close()


def test_i8_recipe_that_fakes_output_and_exits_is_caught(tmp_path: Path) -> None:
    source = recipes.standard({"affiliate_id": "Affiliate ID", "affiliate_name": "Affiliate Name"}, "clean", 1)
    evil = tmp_path / "evil.py"
    evil.write_text(source + "\nprint('{\"ok\": true}')\nraise SystemExit(0)\n")
    result = recipes.check(evil, FIXTURE_DIR / "clean.csv")
    assert not result.ok


@pytest.mark.parametrize("line", ["from onboarding_sdk.recipes import sys", "from onboarding_sdk.read import builtins"])
def test_i8_allow_list_blocks_reexported_modules(tmp_path: Path, line: str) -> None:
    source = (
        line + "\n" + recipes.standard({"affiliate_id": "Affiliate ID", "affiliate_name": "Affiliate Name"}, "clean", 1)
    )
    assert recipes.static_violations(source)


@pytest.mark.parametrize("sponsor", ["x/../sponsor-b", "Sponsor-A", " sponsor-a", "a" * 70])
def test_sponsor_ids_are_validated(tmp_path: Path, sponsor: str) -> None:
    bench = Workbench(offline_services(tmp_path, Models()))
    with pytest.raises(UploadRejected, match="sponsor"):
        bench.create(sponsor_id=sponsor, entity="affiliate", file_name="a.csv", data=b"x", actor=A)


def test_review_workbook_does_not_run_source_formulas(tmp_path: Path) -> None:
    from onboarding_sdk import review
    from onboarding_sdk.canonical import from_table
    from onboarding_sdk.entities.affiliate import process

    path = tmp_path / "f.csv"
    path.write_text('Affiliate ID,Affiliate Name\nAFF_1,"=HYPERLINK(""http://x"",""y"")"\n', encoding="utf-8")
    table = read.open(path).sheets[0].table(1)
    canon = from_table(table, id_column="Affiliate ID", name_column="Affiliate Name")
    data = review.workbook(
        path,
        canon,
        process(canon),
        [{"seq": 1, "kind": "k", "actor": "=cmd()", "at": "t"}],
        {},
        bindings={"affiliate_id": "Affiliate ID", "affiliate_name": "Affiliate Name"},
        header_row=1,
    )
    wb = load_workbook(io.BytesIO(data))
    assert wb["Source"]["B2"].data_type == "s"
    assert wb["Decisions"]["D2"].data_type == "s"
    assert wb["Upload preview"]["B2"].value.startswith("=LEFT(")  # our own formula stays live
