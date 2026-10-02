"""Generate the affiliate test fixtures deterministically.

Usage: uv run python scripts/generate_fixtures.py [--sample-dir DIR] [--out DIR]

Rerunning produces the same bytes: workbook metadata timestamps and zip entry
timestamps are pinned.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import os
import re
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SAMPLES = (
    Path(os.environ.get("STRING_MATCHER_PATH", REPO_ROOT / "../../advance_research/string_matcher_v1"))
    / "docs/sample_data/affiliate"
)
DEFAULT_OUT = REPO_ROOT / "tests/fixtures/affiliate"
FIXED_TIME = datetime(2025, 1, 1, 0, 0, 0)
ZIP_TIME = (1980, 1, 1, 0, 0, 0)
TEMPLATE = ("ITEM_ID", "NAME", "ITEM_TYPE", "DESCRIPTION", "DONOTIMPORT")

# A different batch of synthetic affiliates for the returning-sponsor upload.
RETURNING_ROWS = [
    ("AFF_9101", "Harbor Point GP II, LLC"),
    ("AFF_9102", "Northwind Fund Management, L.P."),
    ("AFF_9103", "Juniper Coinvestment Vehicle, LLC"),
    ("AFF_9104", "Harbor Point Carry Partners, L.P."),
]


def _read_csv(path: Path) -> tuple[list[str], list[list[str]]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    return rows[0], rows[1:]


def _save_xlsx(wb: Workbook, path: Path) -> None:
    wb.properties.created = FIXED_TIME
    wb.properties.modified = FIXED_TIME
    wb.properties.creator = "fixture-generator"
    wb.properties.lastModifiedBy = "fixture-generator"
    raw = io.BytesIO()
    wb.save(raw)
    raw.seek(0)
    out = io.BytesIO()
    with zipfile.ZipFile(raw) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in sorted(src.infolist(), key=lambda info: info.filename):
            info = zipfile.ZipInfo(item.filename, date_time=ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            data = src.read(item.filename)
            if item.filename == "docProps/core.xml":
                # openpyxl stamps the save time into dcterms:modified.
                data = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>2025-01-01T00:00:00Z\g<2>",
                    data,
                )
            dst.writestr(info, data)
    path.write_bytes(out.getvalue())


def _expected_rows(pairs: list[tuple[str, str]]) -> list[dict[str, str]]:
    return [
        {
            "ITEM_ID": item_id,
            "NAME": name,
            "ITEM_TYPE": "Inventory",
            "DESCRIPTION": "",
            "DONOTIMPORT": "",
        }
        for item_id, name in pairs
    ]


def _derive(name: str) -> str:
    # Mirrors the documented rule so expectations stay readable; the rules
    # module is tested against these files, not the other way round.
    upper = re.sub(r"\s+", "_", name.strip().upper())
    return re.sub(r"_+", "_", re.sub(r"[^A-Z0-9_]", "", upper)).strip("_")[:30]


def generate(sample_dir: Path, out: Path) -> list[Path]:
    out.mkdir(parents=True, exist_ok=True)
    (out / "expected").mkdir(exist_ok=True)
    written: list[Path] = []
    expected: dict[str, dict[str, object]] = {}
    golden = json.loads((sample_dir / "expected.json").read_text())["files"]

    header, clean = _read_csv(sample_dir / "affiliate.csv")
    clean_pairs = [(row[0], row[1]) for row in clean]

    for name, source in (
        ("clean.csv", "affiliate.csv"),
        ("edge.csv", "affiliate_edge.csv"),
        ("empty.csv", "affiliate_empty.csv"),
        ("extra_columns.csv", "source.csv"),
    ):
        (out / name).write_bytes((sample_dir / source).read_bytes())
        written.append(out / name)

    expected["clean.csv"] = {
        "sheet": "clean",
        "header_row": 1,
        "bindings": {"affiliate_id": "Affiliate ID", "affiliate_name": "Affiliate Name"},
        "questions": 0,
        "rows": _expected_rows(clean_pairs),
        "findings": [],
    }
    expected["edge.csv"] = {
        "sheet": "edge",
        "header_row": 1,
        "bindings": {"affiliate_id": "Affiliate ID", "affiliate_name": "Affiliate Name"},
        "questions": 0,
        "rows": None,
        "findings": [
            {"row": f["row"], "code": f["code"], "severity": f["severity"]}
            for f in golden["affiliate_edge.csv"]["findings"]
            if f["code"] != "AFF_WARN_ITEM_TYPE_OVERRIDE"
        ],
        # What an analyst does in the grid to clear every error.
        "fixes": [
            {"kind": "exclude_row", "row": 1, "reason": "blank row"},
            {"kind": "override_item_id", "row": 2, "value": "AFF_9999"},
            {"kind": "override_item_id", "row": 4, "value": "AFF_9011"},
            {"kind": "override_item_id", "row": 6, "value": "CASCADE_EMP_COINV_BETA"},
        ],
    }
    expected["empty.csv"] = {
        "sheet": "empty",
        "header_row": 1,
        "bindings": {"affiliate_id": "Affiliate ID", "affiliate_name": "Affiliate Name"},
        "questions": 0,
        "rows": [],
        "findings": [{"row": None, "code": "AFF_WARN_ZERO_RECORDS", "severity": "warning"}],
    }
    _, extra = _read_csv(sample_dir / "source.csv")
    extra_pairs = [(row[0] or _derive(row[1]), row[1]) for row in extra]
    expected["extra_columns.csv"] = {
        "sheet": "extra_columns",
        "header_row": 1,
        "bindings": {"affiliate_id": "Affiliate Id", "affiliate_name": "Affiliate Name"},
        "questions": 0,
        "rows": _expected_rows(extra_pairs),
        "findings": [{"row": 3, "code": "AFF_WARN_ITEM_ID_DERIVED", "severity": "warning"}],
    }

    # 5. Titled workbook: title rows, header on row 4, a trailing total row
    # and a notes sheet.
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Affiliates"
    ws.append(["Affiliate Register"])
    ws.append(["Prepared for sponsor-a by the fund administrator"])
    ws.append([])
    ws.append(header)
    for row in clean:
        ws.append(row)
    ws.append(["Total", f"{len(clean)} affiliates", None])
    notes = wb.create_sheet("Notes")
    notes.append(["Notes"])
    notes.append(["Affiliate IDs are assigned by the fund administrator."])
    _save_xlsx(wb, out / "titled.xlsx")
    written.append(out / "titled.xlsx")
    expected["titled.xlsx"] = {
        "sheet": "Affiliates",
        "header_row": 4,
        "bindings": {"affiliate_id": "Affiliate ID", "affiliate_name": "Affiliate Name"},
        "questions": 0,
        "rows": _expected_rows(clean_pairs),
        "findings": [],
        "dropped": [{"row": 13, "reason": "total_row"}],
    }

    # 6. Renamed headers from the matcher's variant corpus plus distractors.
    renamed_header = ["Affiliate Id", "Affi Name", "Zip Code", "Created By", "Notes"]

    def renamed_book(pairs: list[tuple[str, str]], path: Path) -> None:
        book = Workbook()
        sheet = book.active
        assert sheet is not None
        sheet.title = "Export"
        sheet.append(renamed_header)
        for index, (item_id, name) in enumerate(pairs):
            sheet.append([item_id, name, f"{10001 + index}", "analyst", None])
        _save_xlsx(book, path)

    renamed_book(clean_pairs, out / "renamed.xlsx")
    written.append(out / "renamed.xlsx")
    renamed_bindings = {"affiliate_id": "Affiliate Id", "affiliate_name": "Affi Name"}
    expected["renamed.xlsx"] = {
        "sheet": "Export",
        "header_row": 1,
        "bindings": renamed_bindings,
        "questions": 1,
        "rows": _expected_rows(clean_pairs),
        "findings": [],
    }

    # 7. No ID column at all: every ITEM_ID is derived.
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(["Affiliate Name", "Fund Complex"])
    for row in clean:
        writer.writerow([row[1], row[2]])
    (out / "ids_missing.csv").write_text(buffer.getvalue(), encoding="utf-8")
    written.append(out / "ids_missing.csv")
    expected["ids_missing.csv"] = {
        "sheet": "ids_missing",
        "header_row": 1,
        "bindings": {"affiliate_id": None, "affiliate_name": "Affiliate Name"},
        "questions": 0,
        "rows": _expected_rows([(_derive(name), name) for _, name in clean_pairs]),
        "findings": [
            {"row": index, "code": "AFF_WARN_ITEM_ID_DERIVED", "severity": "warning"}
            for index in range(1, len(clean) + 1)
        ],
    }

    # 8. Two candidate sheets: the analyst must say which one is current.
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Affiliates"
    ws.append(header)
    for row in clean:
        ws.append(row)
    old = wb.create_sheet("Affiliates (old)")
    old.append(header)
    for row in clean[:5]:
        old.append([row[0], row[1].replace(", LLC", " LLC"), row[2]])
    _save_xlsx(wb, out / "two_sheets.xlsx")
    written.append(out / "two_sheets.xlsx")
    expected["two_sheets.xlsx"] = {
        "sheet": "Affiliates",
        "header_row": 1,
        "bindings": {"affiliate_id": "Affiliate ID", "affiliate_name": "Affiliate Name"},
        "questions": 1,
        "answers": {"sheet": "Affiliates"},
        "rows": _expected_rows(clean_pairs),
        "findings": [],
    }

    # 9. The same layout as renamed.xlsx with new values.
    renamed_book(RETURNING_ROWS, out / "returning_sponsor.xlsx")
    written.append(out / "returning_sponsor.xlsx")
    expected["returning_sponsor.xlsx"] = {
        "sheet": "Export",
        "header_row": 1,
        "bindings": renamed_bindings,
        "questions": 0,
        "rows": _expected_rows(RETURNING_ROWS),
        "findings": [],
        "after": "renamed.xlsx",
        "model_calls": 0,
    }

    for name, payload in expected.items():
        target = out / "expected" / f"{Path(name).stem}.json"
        target.write_text(json.dumps({"file": name, **payload}, indent=2, sort_keys=True) + "\n")
        written.append(target)
    return written


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sample-dir", type=Path, default=DEFAULT_SAMPLES)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()
    for path in generate(args.sample_dir.resolve(), args.out):
        print(path.relative_to(REPO_ROOT) if path.is_relative_to(REPO_ROOT) else path)


if __name__ == "__main__":
    main()
