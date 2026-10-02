"""Workbook reading: CSV/TSV and XLSX/XLS behind one small interface.

Row numbers are 1-based sheet rows, as a spreadsheet shows them, so lineage
points at the cell an analyst would open.
"""

from __future__ import annotations

import builtins
import csv
import datetime as dt
import io
import re
from collections.abc import Iterator, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal

Value = str | int | float | bool | dt.date | dt.datetime | dt.time | dt.timedelta | None

CSV_SUFFIXES = {".csv", ".tsv", ".txt"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xls"}
_ENCODINGS = ("utf-8-sig", "cp1252", "latin-1")
_TOTAL = re.compile(r"^\s*(grand\s+)?totals?\b", re.IGNORECASE)


class UnsupportedFileError(ValueError):
    """The upload is not a CSV, TSV or Excel workbook."""


def text(value: Value) -> str | None:
    """The cell as trimmed text, or None when blank."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, dt.datetime) and value.time() == dt.time(0):
        value = value.date()
    rendered = value.isoformat() if isinstance(value, dt.date | dt.time) else str(value)
    rendered = rendered.strip()
    return rendered or None


@dataclass(frozen=True, slots=True)
class RowView:
    row_number: int
    values: dict[str, str | None]

    def get(self, column: str | None) -> str | None:
        if column is None:
            return None
        return self.values.get(column)


@dataclass(frozen=True, slots=True)
class Table:
    sheet: str
    header_row: int
    columns: list[str]
    records: list[RowView]
    dropped: list[tuple[int, str]] = field(default_factory=list)

    def rows(self) -> Iterator[RowView]:
        return iter(self.records)


def _column_names(raw: list[Value]) -> list[str]:
    names: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(raw, start=1):
        name = text(value) or f"column_{index}"
        count = seen.get(name, 0) + 1
        seen[name] = count
        names.append(name if count == 1 else f"{name} ({count})")
    return names


@dataclass(frozen=True, slots=True)
class Sheet:
    name: str
    grid: list[list[Value]]

    @property
    def max_row(self) -> int:
        return len(self.grid)

    def cells(self, max_rows: int | None = None) -> list[list[Value]]:
        return self.grid if max_rows is None else self.grid[:max_rows]

    def row(self, row_number: int) -> list[Value]:
        if 1 <= row_number <= len(self.grid):
            return self.grid[row_number - 1]
        return []

    def table(self, header_row: int, *, stop_at_blank: bool = True, drop_total_rows: bool = True) -> Table:
        if not 1 <= header_row <= max(len(self.grid), 1):
            raise ValueError(f"header row {header_row} is outside sheet {self.name!r}")
        header = list(self.row(header_row))
        while header and text(header[-1]) is None:
            header.pop()
        columns = _column_names(header)
        records: list[RowView] = []
        dropped: list[tuple[int, str]] = []
        for row_number in range(header_row + 1, len(self.grid) + 1):
            raw = self.grid[row_number - 1]
            values = [text(v) for v in raw]
            if not any(values):
                if stop_at_blank:
                    break
                continue
            first = next((v for v in values if v), "")
            if drop_total_rows and _TOTAL.match(first):
                dropped.append((row_number, "total_row"))
                continue
            padded = values + [None] * (len(columns) - len(values))
            records.append(RowView(row_number, dict(zip(columns, padded, strict=False))))
        return Table(self.name, header_row, columns, records, dropped)


@dataclass(frozen=True, slots=True)
class Workbook:
    path: Path
    kind: Literal["csv", "excel"]
    sheets: list[Sheet]

    def sheet(self, name: str) -> Sheet:
        for sheet in self.sheets:
            if sheet.name == name:
                return sheet
        raise KeyError(f"no sheet named {name!r}; sheets are {[s.name for s in self.sheets]}")

    def select(self, name: str) -> Sheet:
        """The named sheet; a CSV has one sheet whose name is only its file stem."""
        if self.kind == "csv":
            return self.sheets[0]
        return self.sheet(name)


def _decode(raw: bytes) -> str:
    for encoding in _ENCODINGS:
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnsupportedFileError("could not decode the text file")  # pragma: no cover


def _read_csv(path: Path) -> Workbook:
    content = _decode(path.read_bytes())
    if path.suffix.lower() == ".tsv":
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(content[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows = list(csv.reader(io.StringIO(content, newline=""), delimiter=delimiter))
    grid: list[list[Value]] = [[cell if cell != "" else None for cell in row] for row in rows]
    return Workbook(path, "csv", [Sheet(path.stem, grid)])


def _normalise_row(row: Sequence[object]) -> list[Value]:
    out: list[Value] = []
    for cell in row:
        out.append(None if cell == "" else cell)  # type: ignore[arg-type]
    while out and out[-1] is None:
        out.pop()
    return out


def _trim(grid: list[list[Value]]) -> list[list[Value]]:
    while grid and not grid[-1]:
        grid.pop()
    return grid


def _read_excel(path: Path) -> Workbook:
    try:
        from python_calamine import CalamineWorkbook

        book = CalamineWorkbook.from_path(str(path))
        # skip_empty_area=False keeps leading blank rows and columns, so row
        # numbers are the ones the analyst sees in Excel.
        sheets = [
            Sheet(
                name,
                _trim([_normalise_row(r) for r in book.get_sheet_by_name(name).to_python(skip_empty_area=False)]),
            )
            for name in book.sheet_names
        ]
        return Workbook(path, "excel", sheets)
    except Exception:
        if path.suffix.lower() == ".xls":
            raise
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = [
            Sheet(ws.title, _trim([_normalise_row(list(r)) for r in ws.iter_rows(values_only=True)]))
            for ws in wb.worksheets
        ]
    finally:
        wb.close()
    return Workbook(path, "excel", sheets)


def open(path: str | Path) -> Workbook:
    path = Path(path)
    if not path.is_file():
        raise builtins.FileNotFoundError(path)
    suffix = path.suffix.lower()
    if suffix in CSV_SUFFIXES:
        return _read_csv(path)
    if suffix in EXCEL_SUFFIXES:
        return _read_excel(path)
    raise UnsupportedFileError(f"unsupported file type {suffix!r}; expected CSV, TSV or Excel")
